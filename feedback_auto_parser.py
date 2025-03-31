import tkinter as tk
from tkinter import filedialog, messagebox
import os
import re
from docx import Document
import openpyxl

def parse_feedback(docx_path):
    doc = Document(docx_path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    data = []
    current_class = None
    wonjang_feedback = ""
    buwonjang_feedback = ""
    student_feedbacks = []
    parsing_stage = "start"
    student_pattern = re.compile(r"^([A-Za-z]+)\((.+?)\)")

    for para in paragraphs:
        if para.startswith("Class "):
            if current_class and student_feedbacks:
                for student_name, feedback in student_feedbacks:
                    full_feedback = f"원장님: {wonjang_feedback}\n부원장님: {buwonjang_feedback}\n{student_name}: {feedback}"
                    data.append([current_class, student_name, full_feedback])
            current_class = para.replace("Class ", "").strip()
            wonjang_feedback = ""
            buwonjang_feedback = ""
            student_feedbacks = []
            parsing_stage = "wonjang"
        elif para.startswith("원장님:"):
            wonjang_feedback = para.replace("원장님:", "").strip()
            parsing_stage = "wonjang"
        elif para.startswith("부원장님:"):
            buwonjang_feedback = para.replace("부원장님:", "").strip()
            parsing_stage = "buwonjang"
        elif student_pattern.match(para):
            match = student_pattern.match(para)
            student_name = match.group(1).strip()
            feedback = para[match.end():].strip()
            student_feedbacks.append((student_name, feedback))
            parsing_stage = "student"
        else:
            if parsing_stage == "wonjang":
                wonjang_feedback += " " + para
            elif parsing_stage == "buwonjang":
                buwonjang_feedback += " " + para
            elif parsing_stage == "student" and student_feedbacks:
                student_feedbacks[-1] = (
                    student_feedbacks[-1][0],
                    student_feedbacks[-1][1] + " " + para
                )

    if current_class and student_feedbacks:
        for student_name, feedback in student_feedbacks:
            full_feedback = f"원장님: {wonjang_feedback}\n부원장님: {buwonjang_feedback}\n{student_name}: {feedback}"
            data.append([current_class, student_name, full_feedback])
    return data

def export_to_excel(data, output_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "정리된 피드백"
    ws.append(["Class", "Student", "Feedback"])
    for row in data:
        ws.append(row)
    wb.save(output_path)

def main():
    root = tk.Tk()
    root.withdraw()
    file_path = filedialog.askopenfilename(
        title="피드백 워드 파일 선택",
        filetypes=[("Word files", "*.docx")]
    )
    if not file_path:
        return

    try:
        data = parse_feedback(file_path)
        output_path = os.path.join(os.path.dirname(file_path), "정리된_피드백.xlsx")
        export_to_excel(data, output_path)
        messagebox.showinfo("완료", f"엑셀 파일이 생성되었습니다:\n{output_path}")
    except Exception as e:
        messagebox.showerror("오류 발생", str(e))

if __name__ == "__main__":
    main()

